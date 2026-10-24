import unittest
from io import BytesIO
from zipfile import ZipFile
from xml.sax.saxutils import escape
from docx import Document
from parser import QuestionParser, QuestionType


class ExporterCompatibilityTests(unittest.TestCase):
    def setUp(self):
        self.parser = QuestionParser()

    def parse_docx(self, *paragraphs):
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        return self.parser.parse_from_docx(document)

    def test_word_option_punctuation(self):
        for value, expected in [
            ("A. 甲", "甲"),
            ("B、乙", "乙"),
            ("C．丙", "丙"),
            ("D) 丁", "丁"),
            ("E）戊", "戊"),
            ("Ｆ: 己", "己"),
        ]:
            with self.subTest(value=value):
                match = self.parser._match_option(value)
                self.assertIsNotNone(match)
                self.assertEqual(match.group(2), expected)

    def test_chaoxing_quiz_only_correct_answer_not_my_answer(self):
        questions = self.parse_docx(
            "1 【单选题】",
            "毛泽东思想初步形成于（ ）",
            "",
            "选项：",
            "A. 土地革命战争时期\nB. 抗日战争时期\nC. 解放战争时期",
            "",
            "我的答案：A",
            "答案状态：错误",
            "正确答案：B",
            "得分：0分",
            "答案解析：不应成为题干",
            "",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].content, "毛泽东思想初步形成于（ ）")
        self.assertEqual(questions[0].options[1], "抗日战争时期")
        self.assertEqual(questions[0].answer, "B")
        self.assertEqual(self.parser.stats["imported"], 1)

    def test_explicit_correct_answer_is_required_even_if_status_correct(self):
        questions = self.parse_docx(
            "2 【单选题】",
            "仅有个人作答的题目",
            "A. 甲",
            "B. 乙",
            "我的答案：A",
            "答案状态：正确",
            "",
        )
        self.assertEqual(questions, [])
        self.assertEqual(self.parser.stats["skipped_no_correct_answer"], 1)

    def test_ulearning_only_my_answer_cannot_be_trusted(self):
        questions = self.parse_docx(
            "请从以下选项中选择",
            "A. 甲\nB. 乙",
            "我的答案：B",
            "",
        )
        self.assertEqual(questions, [])
        self.assertEqual(self.parser.stats["skipped_no_correct_answer"], 1)

    def test_yuketang_labeled_answer_text_is_mapped_to_correct_key(self):
        questions = self.parse_docx(
            "1. 单选题",
            "下列正确的是？",
            "",
            "选项：",
            "A. 红色\nB. 蓝色",
            "",
            "我的答案：A. 红色",
            "正确答案：B. 蓝色",
            "",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "B")

    def test_multiple_choice_comma_separator_and_case(self):
        questions = self.parse_docx(
            "2 【多选题】",
            "下列哪些正确？",
            "A、甲\nB、乙\nC、丙",
            "我的答案：B",
            "正确答案：A、C",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "AC")

    def test_correct_fill_blanks_discard_my_answers(self):
        questions = self.parse_docx(
            "3 【填空题】",
            "填入两个内容",
            "我的答案：错；错",
            "正确答案：第1空：北京；第2空：上海",
            "",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "北京；上海")

    def test_true_false_normalized(self):
        questions = self.parse_docx(
            "【判断题】",
            "地球是圆的。",
            "我的答案：错",
            "正确答案：正确",
            "",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "对")

    def test_missing_correct_answer_does_not_bleed_into_next_question(self):
        questions = self.parse_docx(
            "1 【单选题】",
            "第一题不应导入",
            "A. 甲",
            "B. 乙",
            "我的答案：A",
            "",
            "2 【单选题】",
            "第二题有明确正确答案",
            "A. 甲",
            "B. 乙",
            "正确答案：B",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].content, "第二题有明确正确答案")
        self.assertEqual(self.parser.stats["skipped_no_correct_answer"], 1)

    def test_unusable_correct_answer_skipped(self):
        questions = self.parse_docx(
            "3 【单选题】",
            "缺少实际公布答案",
            "A. 甲",
            "B. 乙",
            "正确答案：暂无答案",
            "",
            "4 【单选题】",
            "不能映射的答案",
            "A. 甲",
            "B. 乙",
            "正确答案：其他选项",
            "",
        )
        self.assertEqual(questions, [])
        self.assertEqual(self.parser.stats["skipped_no_correct_answer"], 1)
        self.assertEqual(self.parser.stats["skipped_unusable_answer"], 1)

    def test_full_option_text_matches_key(self):
        questions = self.parse_docx(
            "【单选题】",
            "选什么颜色？",
            "选项：",
            "A. 红色",
            "B. 蓝色",
            "正确答案：蓝色",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "B")

    def test_essay_explicit_reference_answer(self):
        questions = self.parse_docx(
            "5 【简答题】",
            "简述主要思想。",
            "参考答案：坚持原则",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].question_type, QuestionType.ESSAY)
        self.assertEqual(questions[0].answer, "坚持原则")

    def test_manual_breaks_in_word_options(self):
        document = Document()
        p = document.add_paragraph()
        p.add_run("A. 第一项\nB、第二项\nC．第三项")
        self.assertEqual(self.parser._extract_options_from_paragraph(p),
                         ["第一项", "第二项", "第三项"])

    def test_minimal_jszip_word_document_matches_exporter_layout(self):
        # The upstream userscript writes each piece of text as one <w:p>
        # and omits normal python-docx style parts.
        paragraphs = [
            "1 【单选题】",
            "正确的选项是什么？",
            "",
            "选项：",
            "A. 错误",
            "B. 正确",
            "我的答案：A",
            "正确答案：B",
            "",
        ]
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        body = "".join(
            "<w:p><w:r><w:t>" + escape(text) + "</w:t></w:r></w:p>"
            if text else "<w:p/>" for text in paragraphs
        )
        document_xml = (
            '<w:document xmlns:w="' + namespace + '"><w:body>' +
            body + "</w:body></w:document>"
        )
        types_xml = (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"
        )
        relationships_xml = (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>"
        )
        buffer = BytesIO()
        with ZipFile(buffer, "w") as archive:
            archive.writestr("[Content_Types].xml", types_xml)
            archive.writestr("_rels/.rels", relationships_xml)
            archive.writestr("word/document.xml", document_xml)
        buffer.seek(0)
        questions = self.parser.parse_from_docx(Document(buffer))
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "B")
        self.assertEqual(questions[0].options, ["错误", "正确"])

    def test_split_correct_answer_label_and_value(self):
        questions = self.parse_docx(
            "【判断题】",
            "这是正确的判断。",
            "我的答案：错",
            "正确答案：",
            "",
            "对",
            "",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].answer, "对")

    def test_yuketang_type_with_score_metadata(self):
        questions = self.parse_docx(
            "单选题（共5分）",
            "最常见的题型是什么？",
            "选项：",
            "A. 单选题",
            "B. 多选题",
            "正确答案：A. 单选题",
        )
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].question_type, QuestionType.SINGLE_CHOICE)
        self.assertEqual(questions[0].content, "最常见的题型是什么？")
        self.assertEqual(questions[0].answer, "A")

    def test_split_label_never_uses_my_answer_as_correct(self):
        questions = self.parse_docx(
            "【单选题】",
            "这道题应当被跳过",
            "A. 甲",
            "B. 乙",
            "正确答案：",
            "我的答案：A",
        )
        self.assertEqual(questions, [])

    def test_legacy_plain_text_parser(self):
        questions = self.parser.parse([
            "1 【单选题】",
            "下列哪一项？",
            "A. 第一项",
            "B、第二项",
            "正确答案：B",
        ])
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0].options, ["第一项", "第二项"])


if __name__ == "__main__":
    unittest.main()
